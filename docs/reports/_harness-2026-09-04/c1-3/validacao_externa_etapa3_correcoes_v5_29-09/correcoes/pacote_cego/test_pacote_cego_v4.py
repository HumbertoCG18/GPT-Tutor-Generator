"""Contraexemplos da revisão da etapa 3 (PC02–PC31) como testes do pacote cego v4. Escritos ANTES das correções.

Fixtures sintéticas. Documentos completos gerados e reabertos por python-docx, openpyxl, python-pptx e pypdf
(instalados, não declarados no pyproject: só teste); o gerador usa stdlib e PyMuPDF (dependência declarada).
"""
import io
import json
import re
import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))
import gera_pacote_cego as G  # noqa: E402
from test_pacote_cego import carrega, escreve, fixture, pdf_objetos as pdf, pdf_valido  # noqa: E402
from test_pacote_cego_v3 import fixture_inventario, zip_bytes  # noqa: E402

TOKEN = "https://moodle.exemplo/pluginfile.php/1/x?token=" + "c0ffee12" * 4
MOTOR = "computed_unit_slug=U02"


def rezip(dados, trocas):
    """Reescreve membros de um ZIP/OOXML: {nome: função(bytes) -> bytes}, {nome: bytes} (acrescenta) ou {nome: None}
    (remove)."""
    ent, out = zipfile.ZipFile(io.BytesIO(dados)), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for i in ent.infolist():
            if i.filename in trocas and trocas[i.filename] is None:
                continue
            b = ent.read(i)
            f = trocas.get(i.filename)
            z.writestr(i.filename, f(b) if callable(f) else b)
        for nome, b in trocas.items():
            if b is not None and not callable(b) and nome not in ent.namelist():
                z.writestr(nome, b)
    return out.getvalue()


def docx_bytes(texto="Escalonamento de processos", hiperlink=None):
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    d = docx.Document()
    d.add_paragraph(texto)
    if hiperlink:
        d.part.relate_to(hiperlink, RT.HYPERLINK, is_external=True)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def xlsx_bytes(texto="Notas da turma"):
    import openpyxl
    w = openpyxl.Workbook()
    w.active["A1"] = texto
    b = io.BytesIO()
    w.save(b)
    return b.getvalue()


def pptx_bytes(titulo="Aula 1"):
    import pptx
    p = pptx.Presentation()
    p.slides.add_slide(p.slide_layouts[0]).shapes.title.text = titulo
    b = io.BytesIO()
    p.save(b)
    return b.getvalue()


def entidade(b):
    return b.replace(b"computed_unit_slug", b"computed&#95;unit_slug")


def utf16(b):
    return re.sub(r"""encoding=['"]UTF-8['"]""", 'encoding="UTF-16"', b.decode("utf-8"), count=1).encode("utf-16")


def odf(tipo, corpo):
    ns = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
          'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:xlink="http://www.w3.org/1999/xlink"')
    content = f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {ns}><office:body>{corpo}</office:body></office:document-content>'
    manifest = ('<?xml version="1.0" encoding="UTF-8"?><manifest:manifest '
                'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"/>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", f"application/vnd.oasis.opendocument.{tipo}")
        z.writestr("content.xml", content)
        z.writestr("META-INF/manifest.xml", manifest)
    return buf.getvalue()


def pdf_caso(texto="Aula de filas", anexo=None, via="nomes", escapa=False, conteudo=None):
    """PDF de uma página; anexo por árvore de nomes, anotação de anexo ou objeto órfão; nomes opcionalmente escapados."""
    stream = conteudo if conteudo is not None else b"BT /F1 12 Tf 72 700 Td (%s) Tj ET" % texto.encode("latin-1")
    catalogo = b"<< /Type /Catalog /Pages 2 0 R"
    pagina = b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >>"
    objs = [None, b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>", None,
            b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    if anexo is not None:
        objs += [b"<< /Type /Filespec /F (a.txt) /EF << /F 7 0 R >> >>",
                 b"<< /Type /EmbeddedFile /Length %d >>\nstream\n%s\nendstream" % (len(anexo), anexo)]
        if via == "nomes":
            catalogo += b" /Names << /EmbeddedFiles << /Names [(a.txt) 6 0 R] >> >>"
        elif via == "anotacao":
            objs.append(b"<< /Type /Annot /Subtype /FileAttachment /Rect [10 10 30 30] /FS 6 0 R >>")
            pagina += b" /Annots [8 0 R]"
    objs[0], objs[2] = catalogo + b" >>", pagina + b" >>"
    if escapa:
        objs = [o.replace(b"/EmbeddedFiles", b"/Embedded#46iles").replace(b"/EmbeddedFile ", b"/Embedded#46ile ") for o in objs]
    return pdf(objs)


CASOS = {
    # XML: entidade, atributo, UTF-16 (PC02–08, PC27–30)
    "PC27-docx-entidade": ("a.docx", lambda: rezip(docx_bytes(MOTOR), {"word/document.xml": entidade}), "motor"),
    "PC28-xlsx-entidade": ("a.xlsx", lambda: rezip(xlsx_bytes(MOTOR), {"xl/worksheets/sheet1.xml": entidade}), "motor"),
    "PC29-pptx-utf16": ("a.pptx", lambda: rezip(pptx_bytes(MOTOR), {"ppt/slides/slide1.xml": utf16,   # sem o
                        "ppt/printerSettings/printerSettings1.bin": None}), "motor"),   # printerSettings, recusado à parte
    "PC30-docx-hiperlink-token": ("a.docx", lambda: docx_bytes(hiperlink=TOKEN), "credencial"),
    "PC06-odt-atributo": ("a.odt", lambda: odf("text", f'<text:p><text:a xlink:href="{TOKEN}">link</text:a></text:p>'), "credencial"),
    "PC07-ods-atributo": ("a.ods", lambda: odf("spreadsheet", f'<text:p><text:a xlink:href="{TOKEN}">x</text:a></text:p>'), "credencial"),
    "PC08-odp-atributo": ("a.odp", lambda: odf("presentation", f'<text:p><text:a xlink:href="{TOKEN}">x</text:a></text:p>'), "credencial"),
    "PC13-zip-aninhado-docx-atributo": ("a.zip", lambda: zip_bytes({"trab/entrega.docx": docx_bytes(hiperlink=TOKEN)}), "credencial"),
    "docx-dtd-entidade": ("a.docx", lambda: rezip(docx_bytes(), {"word/document.xml": lambda b: b.replace(
        b"?>", b'?><!DOCTYPE d [<!ENTITY x "computed_unit_slug">]>', 1)}), "DTD"),
    # membros não previstos e binários opacos (PC09, PC10)
    "PC09-docx-membro-dat": ("a.docx", lambda: rezip(docx_bytes(), {"word/leitura.dat": MOTOR.encode()}), "desconhecido"),
    "PC10-docx-printerSettings": ("a.docx", lambda: rezip(docx_bytes(), {"word/printerSettings/printerSettings1.bin": b"\x00" + MOTOR.encode()}), "opaco"),
    "docx-fonte-embutida": ("a.docx", lambda: rezip(docx_bytes(), {"word/fonts/font1.odttf": b"\x00\x01\x00\x00"}), "opaco"),
    "docx-emf": ("a.docx", lambda: rezip(docx_bytes(), {"word/media/image9.emf": b"\x01\x00\x00\x00"}), "opaco"),
    "zip-com-doc-ole": ("a.zip", lambda: zip_bytes({"antigo.doc": b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"}), "opaco"),
    # caminhos do Leitor dentro do ZIP (PC11, PC12)
    "PC11-zip-manual-review": ("a.zip", lambda: zip_bytes({"manual-review/leitura.md": "Leitura da semana"}), "gerado"),
    "PC12-zip-course-map": ("a.zip", lambda: zip_bytes({"COURSE_MAP.md": "Mapa do curso"}), "gerado"),
    # texto: UTF-16 e entidade HTML
    "txt-utf16-motor": ("a.txt", lambda: MOTOR.encode("utf-16"), "motor"),
    "html-entidade-motor": ("a.html", lambda: b"<p>computed&#95;unit_slug=U02</p>", "motor"),
    # PDF: estrutural (anexos) separado do textual (PC19–21, PC31)
    "PC19-pdf-anexo-literal": ("a.pdf", lambda: pdf_caso(anexo=MOTOR.encode()), "embutido"),
    "PC20-pdf-anexo-escapado": ("a.pdf", lambda: pdf_caso(anexo=MOTOR.encode(), escapa=True), "embutido"),
    "pdf-anexo-por-anotacao": ("a.pdf", lambda: pdf_caso(anexo=MOTOR.encode(), via="anotacao"), "embutido"),
    "pdf-anexo-orfao": ("a.pdf", lambda: pdf_caso(anexo=MOTOR.encode(), via="orfao"), "embutido"),
    "PC21-pdf-texto-motor": ("a.pdf", lambda: pdf_caso(texto="computed_unit_slug U02"), "motor"),
    "pdf-invalido": ("a.pdf", lambda: b"%PDF-1.4 bytes quaisquer", "ilegível"),
    # controles: legítimos aceitos
    "PC31-pdf-comentario-apos-eof": ("a.pdf", lambda: pdf_caso() + b"% exemplo didatico: /EmbeddedFile\n", None),
    "controle-pdf-limpo": ("a.pdf", lambda: pdf_caso(), None),
    "controle-docx-limpo": ("a.docx", lambda: docx_bytes(), None),
    "controle-xlsx-limpo": ("a.xlsx", lambda: xlsx_bytes(), None),
    "controle-odt-limpo": ("a.odt", lambda: odf("text", "<text:p>Aula</text:p>"), None),
    "controle-zip-limpo": ("a.zip", lambda: zip_bytes({"aula.pdf": pdf_caso(), "lista.txt": "Exercícios"}), None),
    # recusa de material legítimo documentada (decisão conservadora do Gate 1): pptx padrão tem printerSettings
    "pptx-padrao-printerSettings-recusado": ("a.pptx", lambda: pptx_bytes(), "opaco"),
}


@pytest.mark.parametrize("caso", sorted(CASOS))
def test_verifica_material_v4(caso):
    nome, gera, recusa = CASOS[caso]
    motivo = G.verifica_material(nome, gera())
    assert (motivo is None) if recusa is None else (motivo is not None and recusa in motivo), motivo


def test_PC27_plano_docx_com_entidade_nao_gera_pacote(tmp_path):
    fonte = fixture(tmp_path)
    escreve(tmp_path / "plano/plano.docx", rezip(docx_bytes(MOTOR), {"word/document.xml": entidade}))
    fonte["plano_arquivo"] = str(tmp_path / "plano/plano.docx")
    with pytest.raises(G.PlanoRecusado, match="motor"):
        G.monta(fonte, tmp_path / "pacote")


def test_PC26_anexo_de_pagina_na_arvore_nao_vira_sobra(tmp_path):
    fonte = fixture(tmp_path)
    raiz = Path(fonte["moodle_dir"])
    figura = b"\x89PNG\r\n\x1a\nfigura"
    escreve(raiz / "raw/moodle/pages/501-anexos/figura.png", figura)
    escreve(raiz / "stash/Unidade 1/figura.png", figura)
    api = carrega(raiz / "raw/moodle/contents.json")
    api[1]["modules"][2]["contents"] = [{"type": "file", "filename": "figura.png", "filepath": "/", "filesize": len(figura)}]
    escreve(raiz / "raw/moodle/contents.json", json.dumps(api))
    G.monta(fonte, tmp_path / "pacote")
    p = tmp_path / "pacote"
    assert G.audita(fonte, p) == []
    figs = [m for m in carrega(p / "materiais.json")["materiais"] if m["nome_original"] == "figura.png"]
    assert [(m["tipo"], m["adjudicavel"]) for m in figs] == [("anexo_de_pagina", False)]
    assert not any(q.name.endswith("figura.png") for q in (p / "materiais").iterdir())
    nao_usados = {x["fonte"]: x["motivo"] for x in carrega(p / "pacote_manifesto.json")["fontes_nao_usadas"]}
    assert "anexo de página" in nao_usados["moodle:stash/Unidade 1/figura.png"]


def test_PC25_ligacao_sem_inventario_nao_confirma_proveniencia(tmp_path):
    fonte = fixture(tmp_path)
    G.monta(fonte, tmp_path / "sem")
    lidas = carrega(tmp_path / "sem/pacote_manifesto.json")["fontes_lidas"]
    assert lidas["moodle:stash/Geral/aula1.pdf"]["proveniencia_confirmada"] is False
    fonte2, inv = fixture_inventario(tmp_path / "com")
    fonte2.update(inv, inventario_sha256=G.sha_b(Path(inv["inventario"]).read_bytes()))
    G.monta(fonte2, tmp_path / "com/pacote")
    lidas = carrega(tmp_path / "com/pacote/pacote_manifesto.json")["fontes_lidas"]
    assert lidas["moodle:stash/Geral/aula1.pdf"]["proveniencia_confirmada"] is True


@pytest.mark.parametrize("esperado, erro", [(None, "sem sha256 esperado"), ("0" * 64, "sha256 do inventário")],
                         ids=["sem-sha-esperado", "sha-divergente"])
def test_inventario_exige_sha_esperado(tmp_path, esperado, erro):
    fonte, inv = fixture_inventario(tmp_path)
    fonte.update(inv)
    if esperado:
        fonte["inventario_sha256"] = esperado
    with pytest.raises(G.DivergenciaProveniencia, match=erro):
        G.monta(fonte, tmp_path / "pacote")


def test_com_inventario_arquivo_fora_do_inventario_nao_e_entregue(tmp_path):
    fonte, inv = fixture_inventario(tmp_path)
    fonte.update(inv, inventario_sha256=G.sha_b(Path(inv["inventario"]).read_bytes()))
    escreve(Path(fonte["moodle_dir"]) / "stash/Geral/extra.pdf", pdf_valido("extra"))
    G.monta(fonte, tmp_path / "pacote")
    p = tmp_path / "pacote"
    assert G.audita(fonte, p) == []
    assert "extra.pdf" not in {m["nome_original"] for m in carrega(p / "materiais.json")["materiais"]}
    nao_usados = {x["fonte"]: x["motivo"] for x in carrega(p / "pacote_manifesto.json")["fontes_nao_usadas"]}
    assert "fora do inventário" in nao_usados["moodle:stash/Geral/extra.pdf"]


def test_pdf_sem_camada_de_texto_fica_registrado(tmp_path):
    fonte = fixture(tmp_path, formato="pastas")
    escreve(Path(fonte["moodle_dir"]) / "Unidade 1/figura.pdf", pdf_caso(conteudo=b"0 0 m 100 100 l S"))
    G.monta(fonte, tmp_path / "pacote")
    lidas = carrega(tmp_path / "pacote/pacote_manifesto.json")["fontes_lidas"]
    assert lidas["moodle:Unidade 1/figura.pdf"]["paginas_sem_texto"] == 1
    assert lidas["moodle:Unidade 1/aula1.pdf"]["paginas_sem_texto"] == 0


def test_documentos_completos_reabrem_nas_bibliotecas_consumidoras():
    """Garante que os casos XML são documentos reais, não só bytes: as bibliotecas consumidoras leem a saída do motor."""
    import docx
    import openpyxl
    import pptx
    d = docx.Document(io.BytesIO(rezip(docx_bytes(MOTOR), {"word/document.xml": entidade})))
    assert d.paragraphs[-1].text == MOTOR
    w = openpyxl.load_workbook(io.BytesIO(rezip(xlsx_bytes(MOTOR), {"xl/worksheets/sheet1.xml": entidade})))
    assert w.active["A1"].value == MOTOR
    p = pptx.Presentation(io.BytesIO(rezip(pptx_bytes(MOTOR), {"ppt/slides/slide1.xml": utf16})))
    assert p.slides[0].shapes.title.text == MOTOR
    import pypdf
    r = pypdf.PdfReader(io.BytesIO(pdf_caso(anexo=MOTOR.encode(), escapa=True)))
    assert list(r.attachments) == ["a.txt"] and r.attachments["a.txt"][0] == MOTOR.encode()
